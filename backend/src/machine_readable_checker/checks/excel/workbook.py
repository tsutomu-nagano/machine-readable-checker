from __future__ import annotations

import struct
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
import xlrd

from ...findings import CheckResult, add_finding
from ..common.table import check_rows


def check_xlsx_file(path: Path) -> CheckResult:
    result = CheckResult(str(path))
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
        if not workbook.worksheets:
            add_finding(result, "empty-workbook", "ワークシートがありません。", "error")
            return result
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            _check_xlsx_merged_cells(result, sheet)
            _check_xlsx_objects(result, sheet)
            _check_xlsx_formulas(result, sheet)
            rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
            result.sheet_previews.append(_build_sheet_preview(
                sheet.title,
                rows,
                [
                    {
                        "start_row": merged.min_row,
                        "end_row": merged.max_row,
                        "start_column": merged.min_col,
                        "end_column": merged.max_col,
                    }
                    for merged in sheet.merged_cells.ranges
                ],
            ))
            sheet_result = check_rows(rows, f"{path}:{sheet.title}")
            result.findings.extend(replace(finding, sheet=sheet.title) for finding in sheet_result.findings)
        result.findings = [_with_xlsx_preview(finding, workbook) for finding in result.findings]
        workbook.close()
    except (OSError, InvalidFileException, ValueError) as error:
        add_finding(result, "invalid-xlsx", f"XLSX を読み取れません: {error}", "error")
    return result


def check_xls_file(path: Path) -> CheckResult:
    result = CheckResult(str(path))
    try:
        workbook = xlrd.open_workbook(path, formatting_info=True)
        if not workbook.nsheets:
            add_finding(result, "empty-workbook", "ワークシートがありません。", "error")
            return result
        for sheet in workbook.sheets():
            if getattr(sheet, "visibility", 0) != 0:
                continue
            _check_xls_merged_cells(result, sheet)
            rows = [[sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)] for row_index in range(sheet.nrows)]
            result.sheet_previews.append(_build_sheet_preview(
                sheet.name,
                rows,
                [
                    {
                        "start_row": row_start + 1,
                        "end_row": row_end,
                        "start_column": column_start + 1,
                        "end_column": column_end,
                    }
                    for row_start, row_end, column_start, column_end in sheet.merged_cells
                ],
            ))
            sheet_result = check_rows(rows, f"{path}:{sheet.name}")
            result.findings.extend(replace(finding, sheet=sheet.name) for finding in sheet_result.findings)
        result.findings = [_with_xls_preview(finding, workbook) for finding in result.findings]
    except (AssertionError, OSError, struct.error, xlrd.XLRDError, ValueError) as error:
        add_finding(result, "invalid-xlsx", f"XLS を読み取れません: {error}", "error")
    return result


def _check_xlsx_merged_cells(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-3
    for merged_range in sheet.merged_cells.ranges:
        merged_values = (
            sheet.cell(row=row_index, column=column_index).value
            for row_index in range(merged_range.min_row, merged_range.max_row + 1)
            for column_index in range(merged_range.min_col, merged_range.max_col + 1)
        )
        if not any(str(value).strip() for value in merged_values if value is not None):
            continue
        add_finding(
            result,
            "merged-cells",
            "セル結合は使用しないでください。",
            sheet=sheet.title,
            row=merged_range.min_row,
            column=merged_range.min_col,
            value=str(merged_range),
        )


def _check_xlsx_objects(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-7
    if sheet._images or sheet._charts:
        add_finding(result, "xlsx-object", "図形・画像等のオブジェクトではなくセルにデータを入力してください。", sheet=sheet.title)


def _check_xlsx_formulas(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-6
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                add_finding(result, "formulas", "結果表は数式ではなく値として出力してください。", sheet=sheet.title, row=cell.row, column=cell.column, value=str(cell.value))


def _check_xls_merged_cells(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-3
    for row_start, row_end, column_start, column_end in sheet.merged_cells:
        merged_values = (
            sheet.cell_value(row_index, column_index)
            for row_index in range(row_start, row_end)
            for column_index in range(column_start, column_end)
        )
        if not any(str(value).strip() for value in merged_values if value not in (None, "")):
            continue
        add_finding(
            result,
            "merged-cells",
            "セル結合は使用しないでください。",
            sheet=sheet.name,
            row=row_start + 1,
            column=column_start + 1,
            value=_format_xls_range(row_start, row_end, column_start, column_end),
        )


def _format_xls_range(row_start: int, row_end: int, column_start: int, column_end: int) -> str:
    start = f"{get_column_letter(column_start + 1)}{row_start + 1}"
    end = f"{get_column_letter(column_end)}{row_end}"
    return f"{start}:{end}"


def _with_xlsx_preview(finding, workbook):
    if not finding.sheet or finding.row is None:
        return finding
    sheet = workbook[finding.sheet]
    preview = _build_preview(
        finding,
        sheet.max_row,
        sheet.max_column,
        lambda row, column: sheet.cell(row=row, column=column).value,
    )
    return replace(finding, preview=preview)


def _with_xls_preview(finding, workbook):
    if not finding.sheet or finding.row is None:
        return finding
    sheet = workbook.sheet_by_name(finding.sheet)
    preview = _build_preview(
        finding,
        sheet.nrows,
        sheet.ncols,
        lambda row, column: sheet.cell_value(row - 1, column - 1),
    )
    return replace(finding, preview=preview)


def _build_preview(finding, max_row: int, max_column: int, value_at) -> dict:
    focus_column = finding.column or 1
    start_row = max(1, finding.row - 2)
    end_row = min(max(max_row, finding.row), finding.row + 2)
    start_column = max(1, focus_column - 2)
    end_column = min(max(max_column, focus_column), focus_column + 3)
    rows = [
        [_preview_value(value_at(row, column)) for column in range(start_column, end_column + 1)]
        for row in range(start_row, end_row + 1)
    ]
    return {
        "sheet": finding.sheet,
        "start_row": start_row,
        "start_column": start_column,
        "columns": [get_column_letter(column) for column in range(start_column, end_column + 1)],
        "rows": rows,
        "focus_row": finding.row,
        "focus_column": finding.column,
    }


def _preview_value(value) -> str:
    if value is None:
        return ""
    return str(value)


def _build_sheet_preview(sheet: str, rows: list[list], merged_ranges: list[dict]) -> dict:
    column_count = max((len(row) for row in rows), default=0)
    return {
        "sheet": sheet,
        "columns": [get_column_letter(column) for column in range(1, column_count + 1)],
        "rows": [
            [_preview_value(row[index]) if index < len(row) else "" for index in range(column_count)]
            for row in rows
        ],
        "merged_ranges": merged_ranges,
    }

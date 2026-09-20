import unittest
import struct
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest.mock import patch

from openpyxl import Workbook
import xlwt

from machine_readable_checker.checker import check_file, check_rows


def codes(rows):
    return {item.code for item in check_rows(rows).findings}


class CheckerTests(unittest.TestCase):
    def test_detects_supported_csv_encodings(self):
        cases = (
            ("utf-8-sig", "utf-8-sig"),
            ("utf-8", "utf-8"),
            ("cp932", "cp932"),
        )
        with TemporaryDirectory() as directory:
            for file_encoding, expected in cases:
                path = Path(directory) / f"table-{expected}.csv"
                path.write_bytes("年,人口（人）\n2025,100\n".encode(file_encoding))

                result = check_file(path)

                self.assertEqual(result.encoding, expected)
                self.assertEqual(result.as_dict()["encoding"], expected)

    def test_valid_table_has_no_findings(self):
        self.assertEqual(codes([["年", "人口（人）"], ["2025", "1234"]]), set())


    def test_checks_headers_and_structure(self):
        result = codes([["年", "年", ""], ["2025", "1"], [], ["2026", "2", "3"]])
        self.assertTrue({"duplicate-header", "missing-header", "inconsistent-columns", "split-table"} <= result)

    def test_title_is_not_treated_as_header_when_numeric_data_is_below(self):
        result = check_rows(
            [
                ["平均年齢・平均給与月額等", "", "", ""],
                [],
                ["団体コード", "指定都市名", "給与", ""],
                ["", "", "平均給料月額", "諸手当月額"],
                [],
                ["011002", "札幌市", "301905", "91545"],
                ["041009", "仙台市", "321623", "118023"],
            ]
        )

        self.assertNotIn("missing-header", {item.code for item in result.findings})

    def test_missing_header_is_detected_relative_to_numeric_data(self):
        result = check_rows(
            [
                ["統計表の表題", "", ""],
                [],
                ["年", "人口", ""],
                ["2025", "100", "200"],
                ["2026", "110", "210"],
            ]
        )

        missing = [item for item in result.findings if item.code == "missing-header"]
        self.assertEqual([(item.row, item.column) for item in missing], [(3, 3)])

    def test_blank_rows_outside_numeric_data_are_not_treated_as_split_table(self):
        result = check_rows(
            [
                ["統計表の表題", "", ""],
                [],
                ["年", "地域", "人口"],
                [],
                ["2025", "札幌市", "100"],
                ["2026", "仙台市", "110"],
                [],
                ["注：人口は推計値。", "", ""],
            ]
        )

        self.assertNotIn("split-table", {item.code for item in result.findings})

    def test_blank_row_inside_numeric_data_is_treated_as_split_table(self):
        result = check_rows(
            [
                ["年", "地域", "人口"],
                ["2025", "札幌市", "100"],
                [],
                ["2026", "仙台市", "110"],
            ]
        )

        split = [item for item in result.findings if item.code == "split-table"]
        self.assertEqual([item.row for item in split], [4])


    def test_checks_decorated_values_and_layout(self):
        result = codes([["年", "人口（人）"], ["令和 7年", "1,200 人"], ["2026", "A  B"]])
        self.assertTrue({"era-only-date", "decorated-number", "layout-whitespace"} <= result)

    def test_checks_missing_units_area_abbreviations_empty_numeric_values_and_multiple_sets(self):
        result = codes(
            [
                ["都道府県", "人口"],
                ["青森", "100"],
                ["岩手県", ""],
                ["宮城県", "***"],
                ["都道府県", "人口"],
                ["福島県", "200"],
            ]
        )
        self.assertTrue({"missing-unit", "area-abbreviation", "ambiguous-empty-value", "multiple-table-sets"} <= result)

    def test_zero_is_not_treated_as_a_special_symbol_for_empty_numeric_values(self):
        result = codes([["年", "人口"], ["2025", "100"], ["2026", "0"], ["2027", ""]])
        self.assertNotIn("ambiguous-empty-value", result)

    def test_reads_xlsx_with_openpyxl_and_detects_workbook_features(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "table.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["年", "人口"])
            sheet.append(["2025", "100"])
            sheet["C1"] = "結合項目"
            sheet.merge_cells("C1:D1")
            sheet["B2"] = "=50+50"
            workbook.save(path)
            workbook.close()

            result = check_file(path)

        self.assertTrue({"merged-cells", "formulas"} <= {item.code for item in result.findings})
        self.assertEqual(result.sheet_previews[0]["merged_ranges"], [{
            "start_row": 1,
            "end_row": 1,
            "start_column": 3,
            "end_column": 4,
        }])

    def test_ignores_empty_merged_cells_in_xlsx(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "empty_merged.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["年", "人口"])
            sheet.append(["2025", "100"])
            sheet.merge_cells("C1:D1")
            workbook.save(path)
            workbook.close()

            result = check_file(path)

        self.assertNotIn("merged-cells", {item.code for item in result.findings})

    def test_xlsx_findings_include_sheet_name(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "multi_sheet.xlsx"
            workbook = Workbook()
            first_sheet = workbook.active
            first_sheet.title = "人口"
            first_sheet.append(["年", "人口"])
            first_sheet.append(["2025", "1,200 人"])
            second_sheet = workbook.create_sheet("世帯")
            second_sheet.append(["年", "世帯数"])
            second_sheet.append(["2025", "100"])
            second_sheet["B2"] = "=50+50"
            workbook.save(path)
            workbook.close()

            result = check_file(path)

        sheet_by_code = {finding.code: finding.sheet for finding in result.findings}
        self.assertEqual(sheet_by_code["decorated-number"], "人口")
        self.assertEqual(sheet_by_code["formulas"], "世帯")
        finding_by_code = {finding.code: finding for finding in result.findings}
        preview = finding_by_code["decorated-number"].preview
        self.assertEqual(preview["sheet"], "人口")
        self.assertEqual(preview["focus_row"], 2)
        self.assertEqual(preview["focus_column"], 2)
        self.assertIn("1,200 人", preview["rows"][1])
        population_sheet = next(item for item in result.sheet_previews if item["sheet"] == "人口")
        self.assertEqual(population_sheet["columns"], ["A", "B"])
        self.assertEqual(population_sheet["rows"], [["年", "人口"], ["2025", "1,200 人"]])

    def test_ignores_hidden_sheets_in_xlsx(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "hidden_sheet.xlsx"
            workbook = Workbook()
            visible_sheet = workbook.active
            visible_sheet.title = "表示"
            visible_sheet.append(["年", "人口（人）"])
            visible_sheet.append(["2025", "100"])
            hidden_sheet = workbook.create_sheet("非表示")
            hidden_sheet.append(["年", "人口"])
            hidden_sheet.append(["令和 7年", "1,200 人"])
            hidden_sheet["C1"] = "結合項目"
            hidden_sheet.merge_cells("C1:D1")
            hidden_sheet["B2"] = "=50+50"
            hidden_sheet.sheet_state = "hidden"
            workbook.save(path)
            workbook.close()

            result = check_file(path)

        self.assertEqual(result.findings, [])

    def test_ignores_very_hidden_sheets_in_xlsx(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "very_hidden_sheet.xlsx"
            workbook = Workbook()
            visible_sheet = workbook.active
            visible_sheet.title = "表示"
            visible_sheet.append(["年", "人口（人）"])
            visible_sheet.append(["2025", "100"])
            hidden_sheet = workbook.create_sheet("非表示")
            hidden_sheet.append(["年", "人口"])
            hidden_sheet.append(["令和 7年", "1,200 人"])
            hidden_sheet.sheet_state = "veryHidden"
            workbook.save(path)
            workbook.close()

            result = check_file(path)

        self.assertEqual(result.findings, [])

    def test_reads_xls_with_xlrd_and_includes_sheet_name(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "table.xls"
            workbook = xlwt.Workbook()
            first_sheet = workbook.add_sheet("人口")
            first_sheet.write(0, 0, "年")
            first_sheet.write(0, 1, "人口")
            first_sheet.write(1, 0, "2025")
            first_sheet.write(1, 1, "1,200 人")
            second_sheet = workbook.add_sheet("世帯")
            second_sheet.write_merge(0, 0, 1, 2, "結合項目")
            second_sheet.write(1, 0, "2025")
            second_sheet.write(1, 1, "100")
            workbook.save(path)

            result = check_file(path)

        sheet_by_code = {finding.code: finding.sheet for finding in result.findings}
        self.assertEqual(sheet_by_code["decorated-number"], "人口")
        self.assertEqual(sheet_by_code["merged-cells"], "世帯")
        preview = next(finding.preview for finding in result.findings if finding.code == "decorated-number")
        self.assertEqual(preview["sheet"], "人口")
        self.assertIn("1,200 人", preview["rows"][1])
        checks = {item["id"]: item["status"] for item in result.as_dict()["checks"]}
        self.assertEqual(checks["estat-2-3"], "issues_found")
        self.assertEqual(checks["estat-2-6"], "unchecked")
        self.assertEqual(checks["estat-4-6"], "not_applicable")
        self.assertEqual(result.as_dict()["summary"]["unchecked"], 1)

    def test_ignores_hidden_sheets_in_xls(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "hidden_sheet.xls"
            workbook = xlwt.Workbook()
            visible_sheet = workbook.add_sheet("表示")
            visible_sheet.write(0, 0, "年")
            visible_sheet.write(0, 1, "人口（人）")
            visible_sheet.write(1, 0, "2025")
            visible_sheet.write(1, 1, "100")
            hidden_sheet = workbook.add_sheet("非表示")
            hidden_sheet.visibility = 1
            hidden_sheet.write(0, 0, "年")
            hidden_sheet.write(0, 1, "人口")
            hidden_sheet.write(1, 0, "令和 7年")
            hidden_sheet.write(1, 1, "1,200 人")
            workbook.save(path)

            result = check_file(path)

        self.assertEqual(result.findings, [])

    def test_xls_reader_assertion_is_reported_as_invalid_file(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "broken.xls"
            path.write_bytes(b"not a valid xls")

        with patch("machine_readable_checker.checks.excel.workbook.xlrd.open_workbook", side_effect=AssertionError()):
                result = check_file(path)

        finding = result.findings[0]
        self.assertEqual(finding.code, "invalid-xlsx")
        self.assertEqual(finding.severity, "error")

    def test_xls_reader_struct_error_is_reported_as_invalid_file(self):
        with TemporaryDirectory() as directory:
            path = Path(directory) / "broken.xls"
            path.write_bytes(b"not a valid xls")

        with patch("machine_readable_checker.checks.excel.workbook.xlrd.open_workbook", side_effect=struct.error("unpack requires a buffer of 2 bytes")):
                result = check_file(path)

        finding = result.findings[0]
        self.assertEqual(finding.code, "invalid-xlsx")
        self.assertEqual(finding.severity, "error")

    def test_findings_include_the_cell_value(self):
        findings = check_rows([["年", "人口"], ["令和 7年", "1,200 人"]]).findings
        values_by_code = {finding.code: finding.value for finding in findings}
        self.assertEqual(values_by_code["era-only-date"], "令和 7年")
        self.assertEqual(values_by_code["decorated-number"], "1,200 人")

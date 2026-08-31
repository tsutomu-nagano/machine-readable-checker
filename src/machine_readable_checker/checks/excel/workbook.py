from __future__ import annotations

import struct
from dataclasses import replace
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
import xlrd

from ...findings import CheckResult, add_finding
from ..common.table import check_rows


def check_xlsx_file(path: Path) -> CheckResult:
    result = CheckResult(str(path))
    try:
        workbook = load_workbook(path, data_only=False, read_only=False)
        if not workbook.worksheets:
            add_finding(result, "empty-workbook", "ワークシートがありません。", "error")
            return result
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            _check_xlsx_merged_cells(result, sheet)
            _check_xlsx_objects(result, sheet)
            _check_xlsx_formulas(result, sheet)
            rows = [[cell.value for cell in row] for row in sheet.iter_rows()]
            sheet_result = check_rows(rows, f"{path}:{sheet.title}")
            result.findings.extend(replace(finding, sheet=sheet.title) for finding in sheet_result.findings)
        workbook.close()
    except (OSError, InvalidFileException, ValueError) as error:
        add_finding(result, "invalid-xlsx", f"XLSX を読み取れません: {error}", "error")
    return result


def check_xls_file(path: Path) -> CheckResult:
    result = CheckResult(str(path))
    try:
        workbook = xlrd.open_workbook(path, formatting_info=True)
        if not workbook.nsheets:
            add_finding(result, "empty-workbook", "ワークシートがありません。", "error")
            return result
        for sheet in workbook.sheets():
            if getattr(sheet, "visibility", 0) != 0:
                continue
            _check_xls_merged_cells(result, sheet)
            rows = [[sheet.cell_value(row_index, column_index) for column_index in range(sheet.ncols)] for row_index in range(sheet.nrows)]
            sheet_result = check_rows(rows, f"{path}:{sheet.name}")
            result.findings.extend(replace(finding, sheet=sheet.name) for finding in sheet_result.findings)
    except (AssertionError, OSError, struct.error, xlrd.XLRDError, ValueError) as error:
        add_finding(result, "invalid-xlsx", f"XLS を読み取れません: {error}", "error")
    return result


def _check_xlsx_merged_cells(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-3
    for merged_range in sheet.merged_cells.ranges:
        merged_values = (
            sheet.cell(row=row_index, column=column_index).value
            for row_index in range(merged_range.min_row, merged_range.max_row + 1)
            for column_index in range(merged_range.min_col, merged_range.max_col + 1)
        )
        if not any(str(value).strip() for value in merged_values if value is not None):
            continue
        add_finding(
            result,
            "merged-cells",
            "セル結合は使用しないでください。",
            sheet=sheet.title,
            row=merged_range.min_row,
            column=merged_range.min_col,
            value=str(merged_range),
        )


def _check_xlsx_objects(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-7
    if sheet._images or sheet._charts:
        add_finding(result, "xlsx-object", "図形・画像等のオブジェクトではなくセルにデータを入力してください。", sheet=sheet.title)


def _check_xlsx_formulas(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-6
    for row in sheet.iter_rows():
        for cell in row:
            if cell.data_type == "f":
                add_finding(result, "formulas", "結果表は数式ではなく値として出力してください。", sheet=sheet.title, row=cell.row, column=cell.column, value=str(cell.value))


def _check_xls_merged_cells(result: CheckResult, sheet) -> None:
    # 対応チェック項目: 2-3
    for row_start, row_end, column_start, column_end in sheet.merged_cells:
        merged_values = (
            sheet.cell_value(row_index, column_index)
            for row_index in range(row_start, row_end)
            for column_index in range(column_start, column_end)
        )
        if not any(str(value).strip() for value in merged_values if value not in (None, "")):
            continue
        add_finding(
            result,
            "merged-cells",
            "セル結合は使用しないでください。",
            sheet=sheet.name,
            row=row_start + 1,
            column=column_start + 1,
            value=_format_xls_range(row_start, row_end, column_start, column_end),
        )


def _format_xls_range(row_start: int, row_end: int, column_start: int, column_end: int) -> str:
    start = f"{get_column_letter(column_start + 1)}{row_start + 1}"
    end = f"{get_column_letter(column_end)}{row_end}"
    return f"{start}:{end}"
